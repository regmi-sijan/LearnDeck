"""Read files and links and split them into small, addressable chunks.

Every reader returns an `Ingested` object: a title, a document kind, an ordered
list of chunks (page, slide, section, function, file ...) and an optional
nested outline that groups chunks into larger sections.
"""
from __future__ import annotations

import ast
import csv
import io
import json
import math
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path, PurePosixPath
from urllib.parse import parse_qs, unquote, urlparse

from . import config


class IngestError(Exception):
    """A problem the user should see (bad file, private link, ...)."""


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------
@dataclass
class Chunk:
    id: str
    title: str
    text: str
    loc: str = ""                 # human label: "Page 3", "Slide 4", "app.py, lines 10-40"
    lang: str | None = None       # programming language for code
    page: int | None = None       # page number inside a preview PDF
    preview: str | None = None    # key of the preview PDF this page belongs to
    image: str | None = None      # stored image file name (shown in the source panel)
    vision: bool = False          # True when text must be read from the image by AI
    parent: str | None = None     # set for finer passages split from a chunk

    def to_dict(self) -> dict:
        return asdict(self)


@dataclass
class Ingested:
    title: str
    kind: str
    chunks: list[Chunk] = field(default_factory=list)
    outline: list[dict] = field(default_factory=list)   # [{title, chunk_ids, children}]
    previews: dict[str, bytes] = field(default_factory=dict)  # key -> PDF bytes
    images: dict[str, bytes] = field(default_factory=dict)    # file name -> bytes
    notes: list[str] = field(default_factory=list)


class Builder:
    def __init__(self) -> None:
        self.chunks: list[Chunk] = []

    def add(self, title: str, text: str, **kw) -> str:
        cid = f"c{len(self.chunks) + 1}"
        self.chunks.append(Chunk(cid, clean_title(title) or cid, text, **kw))
        return cid


def group(title: str, ids: list[str], children: list[dict] | None = None) -> dict:
    return {"title": clean_title(title) or "Section", "chunk_ids": list(ids), "children": children or []}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
CODE_EXT = {
    ".py": "python", ".pyw": "python", ".js": "javascript", ".mjs": "javascript", ".cjs": "javascript",
    ".jsx": "jsx", ".ts": "typescript", ".tsx": "tsx", ".java": "java", ".kt": "kotlin", ".kts": "kotlin",
    ".scala": "scala", ".groovy": "groovy", ".gradle": "groovy", ".c": "c", ".h": "c", ".cpp": "cpp",
    ".cc": "cpp", ".cxx": "cpp", ".hpp": "cpp", ".hh": "cpp", ".cs": "csharp", ".go": "go", ".rs": "rust",
    ".rb": "ruby", ".php": "php", ".swift": "swift", ".m": "objectivec", ".mm": "objectivec", ".r": "r",
    ".jl": "julia", ".lua": "lua", ".pl": "perl", ".pm": "perl", ".sh": "bash", ".bash": "bash",
    ".zsh": "bash", ".fish": "bash", ".ps1": "powershell", ".bat": "batch", ".cmd": "batch", ".sql": "sql",
    ".html": "html", ".htm": "html", ".css": "css", ".scss": "scss", ".sass": "scss", ".less": "less",
    ".vue": "vue", ".svelte": "svelte", ".dart": "dart", ".ex": "elixir", ".exs": "elixir", ".erl": "erlang",
    ".hs": "haskell", ".ml": "ocaml", ".fs": "fsharp", ".clj": "clojure", ".elm": "elm", ".zig": "zig",
    ".nim": "nim", ".sol": "solidity", ".v": "verilog", ".sv": "verilog", ".vhd": "vhdl", ".asm": "asm",
    ".s": "asm", ".f90": "fortran", ".f": "fortran", ".matlab": "matlab", ".proto": "protobuf",
    ".graphql": "graphql", ".gql": "graphql", ".tf": "hcl", ".yaml": "yaml", ".yml": "yaml",
    ".toml": "toml", ".json": "json", ".xml": "xml", ".ini": "ini", ".cfg": "ini", ".conf": "ini",
    ".env": "ini", ".cmake": "cmake", ".mk": "makefile",
}
CODE_NAMES = {"makefile": "makefile", "dockerfile": "dockerfile", "cmakelists.txt": "cmake",
              "gemfile": "ruby", "rakefile": "ruby", "procfile": "ini", "jenkinsfile": "groovy"}
TEXT_EXT = {".txt", ".md", ".markdown", ".mdx", ".rst", ".tex", ".log", ".org", ".adoc", ".text"}
IMAGE_EXT = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".gif": "image/gif", ".webp": "image/webp"}
SKIP_DIRS = {".git", ".hg", ".svn", "node_modules", "__pycache__", ".venv", "venv", "env", "dist",
             "build", "target", ".next", ".nuxt", "vendor", ".idea", ".vscode", ".mypy_cache",
             ".pytest_cache", "coverage", ".gradle", "Pods", "DerivedData", "__MACOSX", ".tox",
             "site-packages", "bower_components", ".cache", "out", "bin", "obj"}
SKIP_FILES = {"package-lock.json", "yarn.lock", "pnpm-lock.yaml", "poetry.lock", "Cargo.lock",
              "composer.lock", "Gemfile.lock", ".DS_Store", "go.sum"}
KIND_LABEL = {"pdf": "PDF", "pptx": "Slides", "docx": "Word document", "xlsx": "Spreadsheet",
              "csv": "Table", "code": "Code", "repo": "Code repository", "notebook": "Notebook",
              "markdown": "Notes", "text": "Text", "web": "Web page", "image": "Image",
              "collection": "Collection"}


def norm_text(s: str) -> str:
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    s = re.sub(r"[ \t]+\n", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def clean_title(s: str | None, limit: int = 90) -> str:
    s = re.sub(r"\s+", " ", (s or "")).strip(" -–—•·:#*\t")
    return s if len(s) <= limit else s[: limit - 1].rstrip() + "…"


def first_line(text: str, limit: int = 70) -> str:
    for line in text.splitlines():
        line = line.strip()
        if len(line) >= 3 and re.search(r"[A-Za-z0-9]", line):
            return clean_title(line, limit)
    return ""


def stem(name: str) -> str:
    return PurePosixPath(name).stem or name


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def looks_binary(data: bytes) -> bool:
    head = data[:8000]
    if b"\x00" in head:
        return True
    if not head:
        return False
    bad = sum(1 for b in head if b < 9 or (13 < b < 32))
    return bad / len(head) > 0.1


def window_split(text: str, target: int = 3000) -> list[str]:
    """Split prose into ~target-size pieces on paragraph boundaries."""
    paras = re.split(r"\n\s*\n", text)
    out, cur = [], ""
    for p in paras:
        if cur and len(cur) + len(p) > target:
            out.append(cur.strip())
            cur = ""
        while len(p) > target * 1.6:  # one huge paragraph: cut on sentence ends
            cut = p.rfind(". ", 0, target)
            cut = cut + 1 if cut > target // 3 else target
            out.append((cur + p[:cut]).strip())
            cur, p = "", p[cut:]
        cur += p + "\n\n"
    if cur.strip():
        out.append(cur.strip())
    return out or [text]


def _pymupdf():
    try:
        import pymupdf  # type: ignore
    except ImportError:  # older releases
        import fitz as pymupdf  # type: ignore
    return pymupdf


# ---------------------------------------------------------------------------
# LibreOffice (optional) for .ppt/.doc/.xls and slide previews
# ---------------------------------------------------------------------------
def soffice_path() -> str | None:
    for cand in ("soffice", "libreoffice"):
        p = shutil.which(cand)
        if p:
            return p
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac if os.path.exists(mac) else None


def soffice_convert(data: bytes, in_ext: str, out_fmt: str) -> bytes | None:
    exe = soffice_path()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as td:
        src = Path(td) / f"input{in_ext}"
        src.write_bytes(data)
        try:
            subprocess.run([exe, "--headless", "--convert-to", out_fmt, "--outdir", td, str(src)],
                           capture_output=True, timeout=240, check=False)
        except Exception:
            return None
        ext = "." + out_fmt.split(":")[0]
        out = Path(td) / f"input{ext}"
        return out.read_bytes() if out.exists() else None


# ---------------------------------------------------------------------------
# Readers
# ---------------------------------------------------------------------------
def read_pdf(data: bytes, name: str) -> Ingested:
    pymupdf = _pymupdf()
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:
        raise IngestError(f"Could not open the PDF ({e}).")
    if doc.needs_pass:
        raise IngestError("This PDF is password-protected. Remove the password and try again.")
    b = Builder()
    page_ids: list[str] = []
    scanned = 0
    for i, page in enumerate(doc, start=1):
        try:
            text = norm_text(page.get_text("text", sort=True))
        except Exception:
            text = ""
        needs_vision = len(text) < 40 and (page.get_images() or page.get_drawings())
        if needs_vision:
            scanned += 1
        head = first_line(text)
        page_ids.append(b.add(f"Page {i}" + (f": {head}" if head else ""), text, loc=f"Page {i}",
                              page=i, preview="main", vision=bool(needs_vision)))
    meta_title = (doc.metadata or {}).get("title") or ""
    if len(meta_title) < 4 or re.search(r"microsoft|\.docx?|\.pptx?|untitled", meta_title, re.I):
        meta_title = ""
    p1 = _pdf_title(doc[0]) if doc.page_count else ""
    title = meta_title or p1 or stem(name)

    outline: list[dict] = []
    try:
        toc = doc.get_toc(simple=True)
    except Exception:
        toc = []
    if toc:
        top_lvl = min(t[0] for t in toc)
        top = [t for t in toc if t[0] == top_lvl and t[2] >= 1]
        if 2 <= len(top) <= 60:
            for j, (_, t, pg) in enumerate(top):
                end = top[j + 1][2] - 1 if j + 1 < len(top) else doc.page_count
                end = max(end, pg)
                ids = page_ids[pg - 1:end]
                if ids:
                    outline.append(group(t, ids))
    ing = Ingested(title=title, kind="pdf", chunks=b.chunks, outline=outline, previews={"main": data})
    if scanned:
        ing.notes.append(f"{scanned} page{'s have' if scanned > 1 else ' has'} no selectable text (scanned or "
                         f"image-only), so the AI will read {'them' if scanned > 1 else 'it'} from the image.")
    return ing


def _pdf_title(page) -> str:
    """The biggest text near the top of page 1 is usually the title."""
    try:
        spans = [sp for blk in page.get_text("dict")["blocks"] for ln in blk.get("lines", [])
                 if abs(ln.get("dir", (1, 0))[1]) < 0.01  # skip rotated margin text (arXiv stamps)
                 for sp in ln.get("spans", []) if sp["text"].strip() and not sp["text"].startswith("arXiv:")]
    except Exception:
        return ""
    if not spans:
        return ""
    sizes = sorted(sp["size"] for sp in spans)
    median, top = sizes[len(sizes) // 2], sizes[-1]
    if top < median * 1.15:
        return ""
    words = [sp["text"].strip() for sp in spans if sp["size"] >= top - 0.6 and sp["bbox"][1] < page.rect.height * 0.5]
    title = clean_title(" ".join(words), 140)
    return title if 4 <= len(title) <= 140 else ""


def _iter_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def read_pptx(data: bytes, name: str) -> Ingested:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f"Could not open the PowerPoint file ({e}).")
    b = Builder()
    images: dict[str, bytes] = {}
    sections: list[tuple[str, list[str]]] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        lines: list[str] = []
        pictures = []
        for sh in _iter_shapes(slide.shapes):
            if title_shape is not None and sh.shape_id == title_shape.shape_id:
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = (p.text or "").strip()
                    if t:
                        lines.append("  " * (p.level or 0) + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                ctitle = ch.chart_title.text_frame.text if ch.has_title else ""
                lines.append(f"[Chart{': ' + ctitle if ctitle else ''}]")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    pictures.append((sh.width * sh.height, sh.image.blob, sh.image.ext))
                except Exception:
                    pass
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        body = "\n".join(lines)
        text = norm_text("\n".join(x for x in [title, body, f"Speaker notes: {notes}" if notes else ""] if x))
        img_name, vision = None, False
        if pictures and len(body) < 60:
            _, blob, ext = max(pictures, key=lambda p: p[0])
            if ext.lower() in ("png", "jpg", "jpeg", "gif", "webp"):
                img_name = f"slide-{i}.{ext.lower()}"
                images[img_name] = blob
                vision = True
        label = title or first_line(body) or f"Slide {i}"
        cid = b.add(f"Slide {i}: {label}", text, loc=f"Slide {i}", image=img_name, vision=vision)
        layout = (slide.slide_layout.name or "").lower() if slide.slide_layout is not None else ""
        if "section" in layout or not sections:
            sections.append((label, [cid]))
        else:
            sections[-1][1].append(cid)
    if not b.chunks:
        raise IngestError("The presentation has no slides.")
    title = (prs.core_properties.title or "").strip() or b.chunks[0].title.split(": ", 1)[-1] or stem(name)
    outline = [group(t, ids) for t, ids in sections] if len(sections) >= 2 else []
    ing = Ingested(title=title, kind="pptx", chunks=b.chunks, outline=outline, images=images)
    pdf = soffice_convert(data, ".pptx", "pdf")
    if pdf:
        ing.previews["main"] = pdf
        for idx, ch in enumerate(ing.chunks, start=1):
            ch.page, ch.preview = idx, "main"
    return ing


def read_docx(data: bytes, name: str) -> Ingested:
    from docx import Document as Docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = Docx(io.BytesIO(data))
    except Exception as e:
        raise IngestError(f"Could not open the Word file ({e}).")
    doc_title = (d.core_properties.title or "").strip()
    sections: list[dict] = [{"title": "", "level": 1, "lines": []}]
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "").lower()
            if style == "title":
                doc_title = doc_title or text
                continue
            m = re.match(r"heading\s*(\d)", style)
            if m and int(m.group(1)) <= 3:
                sections.append({"title": text, "level": int(m.group(1)), "lines": []})
            elif "list" in style:
                sections[-1]["lines"].append("- " + text)
            else:
                sections[-1]["lines"].append(text)
        elif tag == "tbl":
            t = Table(el, d)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    ct = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != ct:  # merged cells repeat
                        cells.append(ct)
                sections[-1]["lines"].append(" | ".join(cells))
    sections = [s for s in sections if s["lines"] or s["title"]]
    if not sections:
        raise IngestError("The Word document appears to be empty.")
    levels = sorted({s["level"] for s in sections if s["title"]})
    split_levels = set(levels[:2]) if levels else set()
    b = Builder()
    outline: list[dict] = []
    merged: list[dict] = []
    for s in sections:  # fold deeper headings into their parent section
        if merged and s["title"] and s["level"] not in split_levels:
            merged[-1]["lines"].append(f"### {s['title']}")
            merged[-1]["lines"].extend(s["lines"])
        else:
            merged.append(s)
    top_level = min(split_levels) if split_levels else 1
    for s in merged:
        text = norm_text("\n".join(s["lines"]))
        title = s["title"] or (doc_title if not b.chunks else "") or first_line(text) or "Introduction"
        pieces = window_split(text, 6000) if len(text) > 9000 else [text]
        ids = [b.add(title if len(pieces) == 1 else f"{title} (part {k})", piece, loc=title)
               for k, piece in enumerate(pieces, start=1)]
        if s["level"] == top_level or not outline:
            outline.append(group(title, ids))
        else:
            parent = outline[-1]
            if not parent["children"]:
                parent["children"].append(group(parent["title"], list(parent["chunk_ids"])))
            parent["children"].append(group(title, ids))
            parent["chunk_ids"].extend(ids)
    title = doc_title or first_line(b.chunks[0].text, 100) or stem(name)
    return Ingested(title=title, kind="docx", chunks=b.chunks, outline=outline if len(outline) >= 2 else [])


def _table_text(rows: list[list[str]], label: str, total_rows: int | None, total_cols: int | None) -> str:
    dims = f" — about {total_rows} rows × {total_cols} columns" if total_rows else ""
    body = "\n".join("\t".join(r).rstrip() for r in rows if any(c.strip() for c in r))
    text = f"{label}{dims}. First rows (tab-separated):\n{body}"
    return text[:15000] + ("\n…(truncated)" if len(text) > 15000 else "")


def read_xlsx(data: bytes, name: str) -> Ingested:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise IngestError(f"Could not open the spreadsheet ({e}).")
    b = Builder()
    for ws in wb.worksheets:
        rows = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= 250:
                break
            rows.append(["" if v is None else str(v) for v in row])
        if not any(any(c.strip() for c in r) for r in rows):
            continue
        b.add(f"Sheet: {ws.title}", _table_text(rows, f"Sheet '{ws.title}'", ws.max_row, ws.max_column),
              loc=f"Sheet {ws.title}")
    if not b.chunks:
        raise IngestError("The spreadsheet appears to be empty.")
    return Ingested(title=stem(name), kind="xlsx", chunks=b.chunks)


def read_csv(text: str, name: str) -> Ingested:
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if name.lower().endswith(".tsv") else csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    b = Builder()
    b.add(stem(name), _table_text([r for r in rows[:250]], f"Table '{name}'", len(rows),
                                  max((len(r) for r in rows[:50]), default=0)), loc=name)
    return Ingested(title=stem(name), kind="csv", chunks=b.chunks)


DEF_PATTERNS = [
    r"^(export\s+)?(default\s+)?(async\s+)?function\b",
    r"^(export\s+)?(default\s+)?(abstract\s+)?(data\s+|sealed\s+|open\s+)?(class|interface|struct|enum|trait|impl|object|record|protocol|extension|namespace|module|type)\b",
    r"^(export\s+)?(const|let|var)\s+\w+\s*=\s*(async\s*)?(\([^)]*\)|\w+)\s*=>",
    r"^(export\s+)?(const|let|var)\s+\w+\s*=\s*(async\s+)?function\b",
    r"^(pub(\([^)]*\))?\s+)?(async\s+)?(fn|mod|macro_rules!)\b",
    r"^func\b", r"^def\b", r"^sub\b", r"^(local\s+)?function\b",
    r"^(public|private|protected|internal|static|final|abstract|override|virtual|synchronized|inline|extern|unsafe|\s)+[\w<>\[\],.?*&\s]+\s+[\w~]+\s*\([^;]*$",
    r"^[A-Za-z_][\w\s\*&<>:,]*\s+\**[A-Za-z_][\w:]*\s*\([^;]*\)\s*(const\s*)?\{?\s*$",
    r"^(CREATE|ALTER)\s+(TABLE|VIEW|FUNCTION|PROCEDURE|INDEX)\b",
    r"^#{1,3}\s", r"^\[[\w.\- ]+\]\s*$",  # markdown-ish headers in configs, ini sections
]
DEF_RE = re.compile("|".join(f"(?:{p})" for p in DEF_PATTERNS), re.I)


def code_blocks(text: str, lang: str) -> list[tuple[int, int, str]]:
    """Return (start_line, end_line, title) blocks, 1-based inclusive."""
    lines = text.splitlines()
    n = len(lines)
    if n == 0:
        return []
    if n <= 80:
        return [(1, n, "")]
    blocks: list[tuple[int, int, str]] = []
    if lang == "python":
        try:
            tree = ast.parse(text)
            spans = []
            for node in tree.body:
                start = min([node.lineno] + [d.lineno for d in getattr(node, "decorator_list", [])])
                end = getattr(node, "end_lineno", node.lineno)
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    spans.append((start, end, f"def {node.name}()"))
                elif isinstance(node, ast.ClassDef):
                    spans.append((start, end, f"class {node.name}"))
                else:
                    spans.append((start, end, ""))
            # merge runs of plain statements and attach gaps (comments) to the following block
            prev_end = 0
            for start, end, title in spans:
                start = prev_end + 1
                if blocks and not title and not blocks[-1][2]:
                    blocks[-1] = (blocks[-1][0], end, "")
                else:
                    blocks.append((start, end, title))
                prev_end = end
            if blocks and prev_end < n:
                blocks[-1] = (blocks[-1][0], n, blocks[-1][2])
            blocks = [(s, e, t or ("Imports & setup" if s == 1 else "Top-level code")) for s, e, t in blocks]
        except SyntaxError:
            blocks = []
    if not blocks:
        starts = [i for i, ln in enumerate(lines) if ln and not ln[0].isspace() and DEF_RE.match(ln)
                  and not ln.rstrip().endswith(";")]
        if not starts or starts[0] != 0:
            starts = [0] + starts
        for j, s in enumerate(starts):
            e = (starts[j + 1] - 1) if j + 1 < len(starts) else n - 1
            title = clean_title(lines[s].strip().rstrip("{(:").strip(), 60) \
                if j or DEF_RE.match(lines[s] or "") else "Header"
            blocks.append((s + 1, e + 1, title))
    # merge runs of small blocks (≈40-120 lines per chunk), split huge ones
    generic = ("Header", "Imports & setup", "Top-level code")
    acc: list[list] = []  # [start, end, [titles]]
    for s, e, t in blocks:
        if acc and (acc[-1][1] - acc[-1][0] < 40 or e - s < 12) and (e - acc[-1][0]) <= 120:
            acc[-1][1] = e
            acc[-1][2].append(t)
        else:
            acc.append([s, e, [t]])
    merged: list[tuple[int, int, str]] = []
    for s, e, titles in acc:
        named = [t for t in titles if t and t not in generic] or [t for t in titles if t] or [""]
        title = named[0] if len(named) == 1 else f"{named[0]} … {named[-1]}" if len(named) == 2 \
            else f"{named[0]} … {named[-1]} ({len(named)} parts)"
        merged.append((s, e, title))
    out: list[tuple[int, int, str]] = []
    for s, e, t in merged:
        if e - s > 260:
            k = math.ceil((e - s + 1) / 180)
            step = math.ceil((e - s + 1) / k)
            for part in range(k):
                ps, pe = s + part * step, min(e, s + (part + 1) * step - 1)
                out.append((ps, pe, f"{t} (part {part + 1})"))
        else:
            out.append((s, e, t))
    return out


def add_code(b: Builder, text: str, path: str, lang: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\t", "    ")
    lines = text.split("\n")
    ids = []
    blocks = code_blocks(text, lang)
    for s, e, t in blocks:
        snippet = "\n".join(lines[s - 1:e])
        if not snippet.strip():
            continue
        label = f"{path}, lines {s}–{e}"
        title = f"{path}: {t}" if (t and len(blocks) > 1) else path
        ids.append(b.add(title, snippet, loc=label, lang=lang))
    return ids


def read_code(text: str, name: str, lang: str) -> Ingested:
    b = Builder()
    ids = add_code(b, text, name, lang)
    if not ids:
        raise IngestError("The code file is empty.")
    return Ingested(title=PurePosixPath(name).name, kind="code", chunks=b.chunks)


HEADING_RE = re.compile(r"^(#{1,3})\s+(.+?)\s*#*\s*$")
TEX_RE = re.compile(r"^\\(chapter|section|subsection)\*?\{(.+?)\}")


def read_markdown(text: str, name: str, kind: str = "markdown", title: str | None = None) -> Ingested:
    text = norm_text(text)
    if not text:
        raise IngestError("The document is empty.")
    lines = text.split("\n")
    heads: list[tuple[int, int, str]] = []  # (line, level, title)
    in_fence = False
    for i, ln in enumerate(lines):
        if ln.lstrip().startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        m = HEADING_RE.match(ln)
        if m:
            heads.append((i, len(m.group(1)), m.group(2)))
            continue
        m = TEX_RE.match(ln.strip())
        if m:
            heads.append((i, {"chapter": 1, "section": 1, "subsection": 2}[m.group(1)], m.group(2)))
    b = Builder()
    outline: list[dict] = []
    doc_title = title
    if len(heads) >= 2:
        levels = sorted({h[1] for h in heads})
        if len(levels) > 1 and sum(1 for h in heads if h[1] == levels[0]) == 1 and heads[0][1] == levels[0]:
            doc_title = doc_title or heads[0][2]   # single H1 = document title
            levels = levels[1:]
        split = set(levels[:2])
        top = levels[0]
        cuts = [h for h in heads if h[1] in split]
        if cuts[0][0] > 0:
            intro = norm_text("\n".join(lines[: cuts[0][0]]))
            if len(re.sub(r"^#.*$", "", intro, flags=re.M).strip()) > 30:
                outline.append(group("Introduction", [b.add("Introduction", intro, loc="Introduction")]))
        for j, (ln, lvl, t) in enumerate(cuts):
            end = cuts[j + 1][0] if j + 1 < len(cuts) else len(lines)
            body = norm_text("\n".join(lines[ln:end]))
            pieces = window_split(body, 6000) if len(body) > 9000 else [body]
            ids = [b.add(t if len(pieces) == 1 else f"{t} (part {k})", p, loc=clean_title(t, 60))
                   for k, p in enumerate(pieces, start=1)]
            if lvl == top or not outline:
                outline.append(group(t, ids))
            else:
                par = outline[-1]
                if not par["children"]:
                    par["children"].append(group(par["title"], list(par["chunk_ids"])))
                par["children"].append(group(t, ids))
                par["chunk_ids"].extend(ids)
    else:
        if heads:
            doc_title = doc_title or heads[0][2]
        for k, piece in enumerate(window_split(text, 2500), start=1):
            head = first_line(piece, 60)
            b.add(f"Part {k}" + (f": {head}" if head else ""), piece, loc=f"Part {k}")
    return Ingested(title=doc_title or first_line(text, 90) or stem(name), kind=kind, chunks=b.chunks,
                    outline=outline if len(outline) >= 2 else [])


def read_notebook(text: str, name: str) -> Ingested:
    try:
        nb = json.loads(text)
    except json.JSONDecodeError as e:
        raise IngestError(f"Could not read the notebook ({e}).")
    lang = ((nb.get("metadata") or {}).get("kernelspec") or {}).get("language") or "python"
    groups: list[dict] = []
    for cell in nb.get("cells", []):
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else str(src)
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            m = re.match(r"^\s*#{1,3}\s+(.+)", src)
            if m or not groups:
                groups.append({"title": m.group(1) if m else "Introduction", "parts": []})
            groups[-1]["parts"].append(src.strip())
        elif cell.get("cell_type") == "code":
            if not groups:
                groups.append({"title": "Setup", "parts": []})
            part = f"```{lang}\n{src.rstrip()}\n```"
            outs = []
            for o in cell.get("outputs", []) or []:
                t = o.get("text") or (o.get("data") or {}).get("text/plain") or ""
                t = "".join(t) if isinstance(t, list) else str(t)
                if t.strip():
                    outs.append(t.strip()[:1500])
            if outs:
                part += "\nOutput:\n```\n" + "\n".join(outs)[:2000] + "\n```"
            groups[-1]["parts"].append(part)
    if not groups:
        raise IngestError("The notebook has no cells.")
    b = Builder()
    for g in groups:
        b.add(g["title"], "\n\n".join(g["parts"]), loc=f"Cells: {clean_title(g['title'], 40)}", lang=lang)
    title = groups[0]["title"] if groups[0]["title"] not in ("Introduction", "Setup") else stem(name)
    return Ingested(title=title, kind="notebook", chunks=b.chunks)


def html_to_markdown(html: str, url: str | None = None) -> tuple[str, str]:
    title, md = "", ""
    try:
        import trafilatura  # type: ignore

        md = trafilatura.extract(html, url=url, output_format="markdown", include_tables=True,
                                 include_links=False, include_comments=False, favor_recall=True) or ""
        meta = trafilatura.extract_metadata(html)
        title = (meta.title if meta else "") or ""
    except Exception:
        md = ""
    if len(md) < 300:
        from bs4 import BeautifulSoup

        try:
            soup = BeautifulSoup(html, "lxml")
        except Exception:
            soup = BeautifulSoup(html, "html.parser")
        title = title or (soup.title.get_text(strip=True) if soup.title else "")
        for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg"]):
            tag.decompose()
        out = []
        root = soup.find("main") or soup.find("article") or soup.body or soup
        for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote", "tr"]):
            t = el.get_text(" ", strip=True)
            if not t:
                continue
            if el.name in ("h1", "h2", "h3", "h4"):
                out.append("#" * int(el.name[1]) + " " + t)
            elif el.name == "li":
                out.append("- " + t)
            elif el.name == "pre":
                out.append("```\n" + el.get_text() + "\n```")
            elif el.name == "tr":
                out.append(" | ".join(c.get_text(" ", strip=True) for c in el.find_all(["td", "th"])))
            else:
                out.append(t)
        alt = "\n\n".join(out)
        if len(alt) > len(md):
            md = alt
    return clean_title(title, 120), md


def read_html(text: str, name: str, url: str | None = None) -> Ingested:
    title, md = html_to_markdown(text, url)
    if len(md.strip()) < 40:
        raise IngestError("Couldn't find readable text on that page (it may need a login or JavaScript).")
    ing = read_markdown(md, name, kind="web", title=title or None)
    return ing


def read_image(data: bytes, name: str) -> Ingested:
    ext = PurePosixPath(name).suffix.lower() or ".png"
    img_name = f"image{ext}"
    b = Builder()
    b.add(stem(name), "", loc="Image", image=img_name, vision=True)
    return Ingested(title=stem(name), kind="image", chunks=b.chunks, images={img_name: data})


# ---------------------------------------------------------------------------
# Archives and multi-file collections
# ---------------------------------------------------------------------------
def merge(items: list[tuple[str, Ingested]], title: str, kind: str = "collection") -> Ingested:
    """Combine several ingested files into one document, re-numbering chunk ids."""
    out = Ingested(title=title, kind=kind)
    b = Builder()
    for fi, (fname, ing) in enumerate(items, start=1):
        mapping: dict[str, str] = {}
        for ch in ing.chunks:
            nid = f"c{len(b.chunks) + 1}"
            mapping[ch.id] = nid
            new = Chunk(**{**ch.to_dict(), "id": nid})
            if new.preview:
                new.preview = f"f{fi}-{new.preview}"
            if new.image:
                new.image = f"f{fi}-{new.image}"
            if len(items) > 1 and fname and not new.loc.startswith(fname):
                new.loc = f"{fname}, {new.loc[0].lower() + new.loc[1:] if new.loc[:5] in ('Page ', 'Slide', 'Part ', 'Sheet') else new.loc}" if new.loc and new.loc != fname else fname
            b.chunks.append(new)
        for k, v in ing.previews.items():
            out.previews[f"f{fi}-{k}"] = v
        for k, v in ing.images.items():
            out.images[f"f{fi}-{k}"] = v

        def remap(g: dict) -> dict:
            return group(g["title"], [mapping[c] for c in g["chunk_ids"] if c in mapping],
                         [remap(c) for c in g.get("children", [])])

        ids = [mapping[c.id] for c in ing.chunks]
        children = [remap(g) for g in ing.outline]
        out.outline.append(group(fname or ing.title, ids, children))
        out.notes.extend(ing.notes)
    out.chunks = b.chunks
    return out


def read_archive(data: bytes, name: str, subpath: str | None = None) -> Ingested:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise IngestError("The archive is damaged or not a zip file.")
    infos = [i for i in zf.infolist() if not i.is_dir()]
    names = [i.filename for i in infos]
    roots = {n.split("/", 1)[0] for n in names}
    strip = (roots.pop() + "/") if len(roots) == 1 and all("/" in n for n in names) else ""
    sub = (subpath.strip("/") + "/") if subpath else ""
    files: list[tuple[str, bytes]] = []
    skipped_big = 0
    total = 0
    for info in sorted(infos, key=lambda i: i.filename):
        rel = info.filename[len(strip):]
        if sub and not rel.startswith(sub):
            continue
        rel_disp = rel[len(sub):] if sub else rel
        parts = rel.split("/")
        if any(p in SKIP_DIRS or (p.startswith(".") and p not in (".github",)) for p in parts[:-1]):
            continue
        if parts[-1] in SKIP_FILES or parts[-1].startswith("._"):
            continue
        ext = PurePosixPath(rel).suffix.lower()
        doc_like = ext in (".pdf", ".pptx", ".docx", ".xlsx", ".ipynb") or ext in IMAGE_EXT
        limit = 40_000_000 if doc_like else 400_000
        if info.file_size > limit:
            skipped_big += 1
            continue
        if ext in IMAGE_EXT and len(files) > 30:
            continue
        if len(files) >= config.READ_REPO_MAX_FILES or total > config.READ_MAX_CHARS:
            break
        raw = zf.read(info)
        if not doc_like and (looks_binary(raw) or ext in (".zip", ".jar", ".min.js")):
            continue
        if rel_disp.endswith(".min.js") or rel_disp.endswith(".min.css") or rel_disp.endswith(".map"):
            continue
        files.append((rel_disp, raw))
        total += len(raw)
    if not files:
        raise IngestError("No readable files were found in the archive.")
    ing = read_files(files, title=stem(name), as_repo=True)
    if skipped_big:
        ing.notes.append(f"Skipped {skipped_big} very large file(s).")
    return ing


def _readme_rank(path: str) -> tuple[int, int, str]:
    low = path.lower()
    return (0 if PurePosixPath(low).name.startswith("readme") and "/" not in low else 1, low.count("/"), low)


def read_files(files: list[tuple[str, bytes]], title: str | None = None, as_repo: bool = False) -> Ingested:
    """Read many files (an upload batch or an archive) into one document."""
    if len(files) == 1 and not as_repo:
        return read_one(*files[0])
    tops = {n.split("/", 1)[0] for n, _ in files}
    if len(tops) == 1 and all("/" in n for n, _ in files):  # a dropped folder: strip its name
        folder = tops.pop()
        title = title or folder
        files = [(n[len(folder) + 1:], d) for n, d in files]
    code_like = sum(1 for n, _ in files if detect_kind(n, b"") in ("code", "notebook"))
    repo = as_repo and code_like >= max(1, len(files) / 2)
    files = sorted(files, key=lambda f: _readme_rank(f[0]))
    items: list[tuple[str, Ingested]] = []
    errors = []
    for fname, data in files:
        try:
            items.append((fname, read_one(fname, data)))
        except IngestError as e:
            errors.append(f"{fname}: {e}")
    if not items:
        raise IngestError("None of the files could be read. " + " ".join(errors[:3]))
    if repo:
        # one map chunk, then group files by top-level directory
        tree_lines = [f"{n}  ({len(d):,} bytes)" for n, d in files][:400]
        map_ing = Ingested(title="Repository map", kind="text",
                           chunks=[Chunk("c1", "Repository map", "Files in this repository:\n" + "\n".join(tree_lines),
                                         loc="File list")])
        merged = merge([("Repository map", map_ing)] + items, title or "Repository", kind="repo")
        dirs: dict[str, list[dict]] = {}
        order: list[str] = []
        for g in merged.outline:
            top = g["title"].split("/", 1)[0] if "/" in g["title"] else "(root files)"
            if g["title"] == "Repository map":
                top = "Repository map"
            if top not in dirs:
                dirs[top] = []
                order.append(top)
            dirs[top].append(g)
        outline = []
        for top in order:
            gs = dirs[top]
            if top == "(root files)" and (len(gs) <= 6 or len(order) <= 2):
                outline.extend(gs)  # few top-level files: list them directly
                continue
            ids = [c for g in gs for c in g["chunk_ids"]]
            outline.append(gs[0] if len(gs) == 1 else group(top + ("/" if top != "(root files)" else ""), ids, gs))
        merged.outline = outline
    else:
        merged = merge(items, title or (items[0][1].title if len(items) == 1 else f"{len(items)} files"))
    if not title:  # an upload of loose files: name the deck after them
        names = [PurePosixPath(n).name for n, _ in items]
        merged.title = (", ".join(names[:-1]) + " and " + names[-1]) if len(names) <= 3 else \
            f"{names[0]} and {len(names) - 1} more files"
    if errors:
        merged.notes.append("Some files were skipped: " + "; ".join(errors[:5]))
    return merged


# ---------------------------------------------------------------------------
# Detection and dispatch
# ---------------------------------------------------------------------------
def detect_kind(name: str, data: bytes, content_type: str | None = None) -> str:
    p = PurePosixPath(name.lower())
    ext = p.suffix
    ct = (content_type or "").split(";")[0].strip().lower()
    if data[:4] == b"%PDF" or (not data and ext == ".pdf") or ct == "application/pdf":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            zn = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            zn = set()
        if "ppt/presentation.xml" in zn:
            return "pptx"
        if "word/document.xml" in zn:
            return "docx"
        if "xl/workbook.xml" in zn:
            return "xlsx"
        return "archive"
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" or ext in (".ppt", ".doc", ".xls"):
        return "legacy-office"
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:3] == b"\xff\xd8\xff" or data[:6] in (b"GIF87a", b"GIF89a") \
            or (data[:4] == b"RIFF" and data[8:12] == b"WEBP") or (not data and ext in IMAGE_EXT):
        return "image"
    if not data:  # name-only guess
        if ext in (".pptx",):
            return "pptx"
        if ext == ".docx":
            return "docx"
        if ext in (".xlsx", ".xlsm"):
            return "xlsx"
        if ext == ".zip":
            return "archive"
    if ext == ".ipynb":
        return "notebook"
    if ext in (".csv", ".tsv"):
        return "csv"
    if ext in (".html", ".htm", ".xhtml") or ct in ("text/html", "application/xhtml+xml"):
        return "html"
    if ext in (".md", ".markdown", ".mdx", ".rst", ".tex", ".org", ".adoc"):
        return "markdown"
    if ext in CODE_EXT or p.name in CODE_NAMES:
        return "code"
    if ext in TEXT_EXT:
        return "text"
    if data and looks_binary(data):
        return "unknown"
    if data:
        head = decode_text(data[:2000]).lower()
        if "<html" in head or "<!doctype html" in head:
            return "html"
        if re.search(r"^(def |class |import |from \w+ import|#include|package |function |const |public class)",
                     head, re.M):
            return "code"
    return "text"


def guess_lang(name: str, text: str = "") -> str:
    p = PurePosixPath(name.lower())
    if p.name in CODE_NAMES:
        return CODE_NAMES[p.name]
    if p.suffix in CODE_EXT:
        return CODE_EXT[p.suffix]
    if text.startswith("#!"):
        line = text.split("\n", 1)[0]
        for key, lang in (("python", "python"), ("node", "javascript"), ("bash", "bash"), ("sh", "bash")):
            if key in line:
                return lang
    return "python" if re.search(r"^(def |import |from \w+ import)", text, re.M) else "text"


def read_one(name: str, data: bytes, content_type: str | None = None, url: str | None = None) -> Ingested:
    if not data:
        raise IngestError(f"'{name}' is empty.")
    kind = detect_kind(name, data, content_type)
    if kind == "pdf":
        return read_pdf(data, name)
    if kind == "pptx":
        return read_pptx(data, name)
    if kind == "docx":
        return read_docx(data, name)
    if kind == "xlsx":
        return read_xlsx(data, name)
    if kind == "archive":
        return read_archive(data, name)
    if kind == "image":
        return read_image(data, name)
    if kind == "legacy-office":
        ext = PurePosixPath(name.lower()).suffix or ".ppt"
        target = {".ppt": "pptx", ".doc": "docx", ".xls": "xlsx"}.get(ext, "pdf")
        converted = soffice_convert(data, ext, target)
        if not converted:
            raise IngestError(f"Old Office files ({ext}) need LibreOffice to convert them. Install LibreOffice, "
                              f"or open the file and 'Save As' .{target}, then upload that.")
        return read_one(stem(name) + "." + target, converted)
    if kind == "unknown":
        raise IngestError(f"'{name}' isn't a file type I can read yet.")
    text = decode_text(data)
    if kind == "notebook":
        return read_notebook(text, name)
    if kind == "csv":
        return read_csv(text, name)
    if kind == "html":
        return read_html(text, name, url)
    if kind == "markdown":
        return read_markdown(text, name)
    if kind == "code":
        return read_code(text, name, guess_lang(name, text))
    # plain text: treat as markdown if it has headings, else windows
    ing = read_markdown(text, name, kind="text")
    return ing


# ---------------------------------------------------------------------------
# Links
# ---------------------------------------------------------------------------
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/126.0 Safari/537.36 LearnDeck/1.0")


def fetch(url: str) -> tuple[bytes, str, str, str]:
    """Download a URL. Returns (data, content_type, final_url, filename)."""
    import httpx

    limit = config.MAX_UPLOAD_MB * 1024 * 1024
    try:
        with httpx.Client(follow_redirects=True, timeout=httpx.Timeout(60, connect=15),
                          headers={"User-Agent": UA, "Accept": "*/*"}) as client:
            with client.stream("GET", url) as r:
                if r.status_code in (401, 403):
                    raise IngestError(f"The site refused access (HTTP {r.status_code}). If it needs a login, "
                                      "download the file and upload it instead.")
                if r.status_code == 404:
                    raise IngestError("That link returned 'not found' (HTTP 404).")
                if r.status_code >= 400:
                    raise IngestError(f"The link returned an error (HTTP {r.status_code}).")
                buf = bytearray()
                for part in r.iter_bytes():
                    buf.extend(part)
                    if len(buf) > limit:
                        raise IngestError(f"The file is larger than {config.MAX_UPLOAD_MB} MB.")
                ctype = r.headers.get("content-type", "")
                cd = r.headers.get("content-disposition", "")
                final = str(r.url)
    except IngestError:
        raise
    except httpx.HTTPError as e:
        raise IngestError(f"Couldn't download the link ({e.__class__.__name__}: {e}).")
    fname = ""
    m = re.search(r"filename\*=(?:UTF-8'')?([^;]+)|filename=\"?([^\";]+)\"?", cd, re.I)
    if m:
        fname = unquote((m.group(1) or m.group(2) or "").strip().strip('"'))
    if not fname:
        fname = unquote(PurePosixPath(urlparse(final).path).name) or "page"
    if "html" in ctype and not fname.lower().endswith((".html", ".htm")):
        fname = (fname if fname != "page" else urlparse(final).netloc) + ".html"
    return bytes(buf), ctype, final, fname


def read_url(url: str) -> Ingested:
    url = url.strip()
    if not re.match(r"^[a-z][a-z0-9+.-]*://", url, re.I):
        url = "https://" + url
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.netloc:
        raise IngestError("Please paste a full web link (starting with http:// or https://).")
    host = p.netloc.lower().split(":")[0]
    parts = [x for x in p.path.split("/") if x]

    if host in ("youtube.com", "www.youtube.com", "youtu.be", "m.youtube.com", "vimeo.com"):
        raise IngestError("Video links aren't supported. Paste a transcript, article or document link instead.")

    if host in ("github.com", "www.github.com") and len(parts) >= 2:
        owner, repo = parts[0], parts[1].removesuffix(".git")
        if len(parts) >= 4 and parts[2] == "blob":
            raw = f"https://raw.githubusercontent.com/{owner}/{repo}/{'/'.join(parts[3:])}"
            data, ctype, final, _ = fetch(raw)
            ing = read_one(parts[-1], data, None, raw)
            return ing
        ref = parts[3] if len(parts) >= 4 and parts[2] == "tree" else None
        sub = "/".join(parts[4:]) if ref and len(parts) > 4 else None
        urls = [f"https://github.com/{owner}/{repo}/archive/{ref or 'HEAD'}.zip",
                f"https://api.github.com/repos/{owner}/{repo}/zipball" + (f"/{ref}" if ref else "")]
        data, last_err = None, None
        for zurl in urls:
            try:
                data, _, _, _ = fetch(zurl)
                break
            except IngestError as e:
                last_err = e
        if data is None:
            raise IngestError(f"Couldn't download the GitHub repository ({last_err}). Private repos must be "
                              "downloaded as a .zip and uploaded.")
        ing = read_archive(data, f"{repo}.zip", subpath=sub)
        ing.title = f"{owner}/{repo}" + (f"/{sub}" if sub else "")
        return ing

    if host == "docs.google.com":
        m = re.match(r"/(document|presentation|spreadsheets)/d/([\w-]+)", p.path)
        if m:
            kind, fid = m.groups()
            export = {
                "document": f"https://docs.google.com/document/d/{fid}/export?format=docx",
                "presentation": f"https://docs.google.com/presentation/d/{fid}/export/pptx",
                "spreadsheets": f"https://docs.google.com/spreadsheets/d/{fid}/export?format=xlsx",
            }[kind]
            data, ctype, final, fname = fetch(export)
            if "html" in ctype:
                raise IngestError("This Google file isn't public. Set sharing to 'Anyone with the link', "
                                  "or download it and upload the file.")
            return read_one(fname, data, ctype, final)

    if host == "drive.google.com":
        m = re.search(r"/file/d/([\w-]+)", p.path)
        fid = m.group(1) if m else (parse_qs(p.query).get("id") or [None])[0]
        if fid:
            data, ctype, final, fname = fetch(f"https://drive.google.com/uc?export=download&id={fid}")
            if "html" in ctype and b"<html" in data[:2000].lower():
                raise IngestError("Google Drive didn't return the file (it may be private or too large). "
                                  "Download it and upload it instead.")
            return read_one(fname, data, ctype, final)

    if host.endswith("arxiv.org"):
        m = re.match(r"/(abs|html)/(.+?)/?$", p.path)
        if m:
            url = f"https://arxiv.org/pdf/{m.group(2)}"

    data, ctype, final, fname = fetch(url)
    ing = read_one(fname, data, ctype, final)
    if ing.kind == "web" and (not ing.title or ing.title == stem(fname)):
        ing.title = urlparse(final).netloc
    return ing
